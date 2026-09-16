"""Export analysis results to Excel."""

import io

from analyzer import StudentPerformanceAnalyzer
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _autosize_columns(sheet) -> None:
    for col_cells in sheet.columns:
        max_length = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        sheet.column_dimensions[col_letter].width = min(max_length + 2, 40)


def _write_header_row(sheet, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)


def export_analysis_to_workbook(analyzer: StudentPerformanceAnalyzer) -> Workbook:
    """Build an Excel workbook with analysis sheets."""
    wb = Workbook()

    # --- Student Reports ---
    ws_reports = wb.active
    ws_reports.title = "Student Reports"
    reports = analyzer.analyze_all()
    subjects = sorted({subj for r in reports for subj in r.grades})

    report_headers = [
        "student_id", "name", "average", "letter_grade", "performance_level",
        "passed_subjects", "failed_subjects", "strongest_subject", "weakest_subject",
        *subjects,
    ]
    _write_header_row(ws_reports, report_headers)

    for row_idx, report in enumerate(reports, start=2):
        values = [
            report.student_id, report.name, report.average, report.letter_grade,
            report.performance_level, report.passed_subjects, report.failed_subjects,
            report.strongest_subject, report.weakest_subject,
            *[report.grades.get(s) for s in subjects],
        ]
        for col, value in enumerate(values, start=1):
            ws_reports.cell(row=row_idx, column=col, value=value)

    _autosize_columns(ws_reports)

    # --- Class Summary ---
    ws_summary = wb.create_sheet("Class Summary")
    summary = analyzer.class_summary()
    summary_rows = [
        ("Metric", "Value"),
        ("Total Students", summary.total_students),
        ("Class Average", summary.class_average),
        ("Highest Average", summary.highest_average),
        ("Lowest Average", summary.lowest_average),
        ("Pass Rate (%)", summary.pass_rate),
        ("Top Performers", ", ".join(summary.top_performers)),
        ("At-Risk Students", ", ".join(summary.at_risk_students)),
    ]
    for row_idx, (label, value) in enumerate(summary_rows, start=1):
        ws_summary.cell(row=row_idx, column=1, value=label)
        ws_summary.cell(row=row_idx, column=2, value=value)
        if row_idx == 1:
            ws_summary.cell(row=row_idx, column=1).font = Font(bold=True)
            ws_summary.cell(row=row_idx, column=2).font = Font(bold=True)

    start_row = len(summary_rows) + 2
    ws_summary.cell(row=start_row, column=1, value="Subject").font = Font(bold=True)
    ws_summary.cell(row=start_row, column=2, value="Average").font = Font(bold=True)
    for i, (subject, avg) in enumerate(summary.subject_averages.items(), start=start_row + 1):
        ws_summary.cell(row=i, column=1, value=subject)
        ws_summary.cell(row=i, column=2, value=avg)

    _autosize_columns(ws_summary)

    # --- Subject Statistics ---
    ws_subjects = wb.create_sheet("Subject Statistics")
    subject_headers = ["subject", "mean", "median", "min", "max", "std_dev"]
    _write_header_row(ws_subjects, subject_headers)
    for row_idx, (subject, stats) in enumerate(analyzer.subject_statistics().items(), start=2):
        for col, key in enumerate(subject_headers, start=1):
            value = subject if key == "subject" else stats[key]
            ws_subjects.cell(row=row_idx, column=col, value=value)
    _autosize_columns(ws_subjects)

    # --- Rankings ---
    ws_rankings = wb.create_sheet("Rankings")
    ranking_headers = [
        "rank", "student_id", "name", "average", "letter_grade", "performance_level",
    ]
    _write_header_row(ws_rankings, ranking_headers)
    for rank, report in enumerate(analyzer.rank_students(), start=1):
        row_idx = rank + 1
        values = [rank, report.student_id, report.name, report.average,
                  report.letter_grade, report.performance_level]
        for col, value in enumerate(values, start=1):
            ws_rankings.cell(row=row_idx, column=col, value=value)
    _autosize_columns(ws_rankings)

    return wb


def export_analysis_to_bytes(analyzer: StudentPerformanceAnalyzer) -> bytes:
    """Return Excel file contents as bytes for download."""
    buffer = io.BytesIO()
    wb = export_analysis_to_workbook(analyzer)
    wb.save(buffer)
    return buffer.getvalue()


def export_analysis_to_file(analyzer: StudentPerformanceAnalyzer, filepath: str) -> None:
    """Save analysis to an .xlsx file."""
    wb = export_analysis_to_workbook(analyzer)
    wb.save(filepath)
