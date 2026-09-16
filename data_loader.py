"""Load student grade data from CSV and Excel files."""

import csv
import io
from pathlib import Path
from typing import BinaryIO, TextIO

from models import Student

REQUIRED_COLUMNS = ("student_id", "name")


def _parse_students(fieldnames: list[str] | None, rows: list[dict[str, str]]) -> list[Student]:
    if not fieldnames or not all(col in fieldnames for col in REQUIRED_COLUMNS):
        raise ValueError("Data must include 'student_id' and 'name' columns")

    subject_columns = [col for col in fieldnames if col not in REQUIRED_COLUMNS]
    students: list[Student] = []

    for row in rows:
        grades: dict[str, float] = {}
        for subject in subject_columns:
            raw = str(row.get(subject, "")).strip()
            if raw:
                try:
                    grades[subject] = float(raw)
                except ValueError as exc:
                    name = row.get("name", "unknown")
                    raise ValueError(f"Invalid grade '{raw}' for {name} in {subject}") from exc

        students.append(
            Student(
                student_id=str(row["student_id"]).strip(),
                name=str(row["name"]).strip(),
                grades=grades,
            )
        )

    if not students:
        raise ValueError("No student records found in the data file")

    return students


def _rows_from_csv(source: TextIO) -> tuple[list[str], list[dict[str, str]]]:
    reader = csv.DictReader(source)
    fieldnames = list(reader.fieldnames or [])
    rows = [{k: v or "" for k, v in row.items()} for row in reader]
    return fieldnames, rows


def _rows_from_excel(source: BinaryIO) -> tuple[list[str], list[dict[str, str]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        raise ValueError("Excel file has no active worksheet")

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ValueError("Excel worksheet is empty") from exc

    fieldnames = [str(cell).strip() if cell is not None else "" for cell in header_row]
    if not any(fieldnames):
        raise ValueError("Excel worksheet has no header row")

    rows: list[dict[str, str]] = []
    for row_values in rows_iter:
        if all(cell is None or str(cell).strip() == "" for cell in row_values):
            continue
        row_dict = {}
        for i, col in enumerate(fieldnames):
            if not col:
                continue
            value = row_values[i] if i < len(row_values) else None
            row_dict[col] = "" if value is None else str(value)
        rows.append(row_dict)

    workbook.close()
    return fieldnames, rows


def load_students(filepath: str | Path) -> list[Student]:
    """Load students from a CSV or Excel file."""
    path = Path(filepath)
    if not path.exists():
        raise FileNotFoundError(f"Data file not found: {path}")

    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8") as f:
            fieldnames, rows = _rows_from_csv(f)
    elif suffix in (".xlsx", ".xlsm"):
        with path.open("rb") as f:
            fieldnames, rows = _rows_from_excel(f)
    else:
        raise ValueError(f"Unsupported file format: {suffix}. Use .csv or .xlsx")

    return _parse_students(fieldnames, rows)


def load_students_from_upload(file_obj: BinaryIO | TextIO, filename: str) -> list[Student]:
    """Load students from an uploaded file object (Streamlit upload)."""
    suffix = Path(filename).suffix.lower()

    if suffix == ".csv":
        if isinstance(file_obj, io.TextIOBase):
            fieldnames, rows = _rows_from_csv(file_obj)
        else:
            text = io.TextIOWrapper(file_obj, encoding="utf-8-sig")
            fieldnames, rows = _rows_from_csv(text)
    elif suffix in (".xlsx", ".xlsm"):
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)
        fieldnames, rows = _rows_from_excel(file_obj)  # type: ignore[arg-type]
    else:
        raise ValueError(f"Unsupported file format: {suffix}. Use .csv or .xlsx")

    return _parse_students(fieldnames, rows)
